from __future__ import annotations

import ast
import io
import re
from typing import Any

from backend.models.scenario_template import OperationalScenario
from openpyxl import load_workbook


class ScenarioExcelParseError(ValueError):
    pass


def _normalize_header(s: Any) -> str:
    """Normalize header string for comparison."""
    return str(s or "").strip().lower()


def parse_scenario_template_excel(file_bytes: bytes) -> list[OperationalScenario]:
    """Parse Excel file into OperationalScenario objects."""
    wb = load_workbook(filename=bytes_to_stream(file_bytes), data_only=True)
    ws = wb.worksheets[0]

    # Get headers once and build index map in single pass
    headers_normalized = []
    name_col = None
    step_cols: list[int] = []

    for col_idx in range(1, ws.max_column + 1):
        header_val = ws.cell(row=1, column=col_idx).value
        normalized = _normalize_header(header_val)
        headers_normalized.append(normalized)

        if normalized in {"name", "scenario", "scenario name", "scenario_name", "title"}:
            name_col = col_idx
        elif normalized.startswith("step") or re.fullmatch(r"s\d+", normalized) or re.fullmatch(r"\d+", normalized):
            step_cols.append(col_idx)

    # Validate required columns
    if name_col is None:
        raise ScenarioExcelParseError("Missing required column: Name")

    if not step_cols:
        step_cols = [col_idx for col_idx in range(1, ws.max_column + 1) if col_idx != name_col]

    if not step_cols:
        raise ScenarioExcelParseError("No step columns found (expected: step 1, step 2, ...)")

    scenarios: list[OperationalScenario] = []
    scenario_index = 0

    # Process data rows
    for row_idx in range(2, ws.max_row + 1):
        name_cell = ws.cell(row=row_idx, column=name_col)
        name_val = name_cell.value

        # Skip empty rows
        if not name_val or not str(name_val).strip():
            continue

        scenario_index += 1
        scenario_name = str(name_val).strip()
        steps: list[dict[str, Any]] = []

        # Extract steps for this scenario
        for col_idx in step_cols:
            step_cell = ws.cell(row=row_idx, column=col_idx)
            cell_val = step_cell.value

            # Skip empty cells
            if not cell_val:
                continue

            txt = str(cell_val).strip()

            # EOE means: stop reading steps for this scenario
            if txt.upper() == "EOE":
                break

            # Parse step as dictionary
            step_data: dict[str, Any]
            if txt.startswith("{") and txt.endswith("}"):
                try:
                    parsed_step = ast.literal_eval(txt)
                except Exception as e:
                    raise ScenarioExcelParseError(
                        f"Invalid step value at row {row_idx}, col {col_idx}: {txt}"
                    ) from e

                if not isinstance(parsed_step, dict):
                    raise ScenarioExcelParseError(
                        f"Step must be a dict at row {row_idx}, col {col_idx}"
                    )

                step_data = parsed_step
            else:
                # Fallback for simpler Excel layouts: treat the cell content as a step type.
                step_data = {"type": txt}

            steps.append(step_data)

        # Create scenario (allows empty steps, as per original behavior)
        scenarios.append(
            OperationalScenario(
                index=scenario_index,
                name=scenario_name,
                steps=steps,
            )
        )

    if not scenarios:
        raise ScenarioExcelParseError("No operational scenarios found in Excel")

    return scenarios


def bytes_to_stream(b: bytes) -> io.BytesIO:
    """Convert bytes to file-like stream object for openpyxl."""
    return io.BytesIO(b)